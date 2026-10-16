#!/usr/bin/env python3
"""docs/TECH_TRANSFER.md를 유지보수 가능한 Word 문서로 변환한다."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent


def plain(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    return text.replace("\\|", "|").strip()


def set_east_asia(run, name: str = "맑은 고딕") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color in ((1, 17, "17365D"), (2, 14, "1F4E78"), (3, 12, "2F5597")):
        style = document.styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "인지검사 문항 생성 파이프라인 기술이전서 · v1.0.0"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_east_asia(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])


def add_table(document: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    for col, value in enumerate(rows[0]):
        cell = table.rows[0].cells[col]
        cell.text = plain(value)
        shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            set_east_asia(run)
            run.bold = True
    for source in rows[1:]:
        cells = table.add_row().cells
        for col in range(width):
            cells[col].text = plain(source[col]) if col < len(source) else ""
            for paragraph in cells[col].paragraphs:
                for run in paragraph.runs:
                    set_east_asia(run)
                    run.font.size = Pt(9)
    document.add_paragraph()


def convert(source: Path, destination: Path) -> None:
    document = Document()
    configure(document)
    document.core_properties.title = "인지검사 문항 생성 파이프라인 기술이전서"
    document.core_properties.subject = "itemgen-bench v1.0.0 기술이전"
    document.core_properties.author = "Lucypothesis contributors"

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    first_title = True
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.style = document.styles["No Spacing"]
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(7)
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            if language:
                run.font.color.rgb = RGBColor(45, 45, 45)
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            title = plain(heading.group(2))
            if first_title and level == 1:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(80)
                paragraph.paragraph_format.space_after = Pt(20)
                run = paragraph.add_run(title)
                set_east_asia(run)
                run.bold = True
                run.font.size = Pt(24)
                run.font.color.rgb = RGBColor(23, 54, 93)
                subtitle = document.add_paragraph("기술이전 납품본 · v1.0.0 · 2026-10-16")
                subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in subtitle.runs:
                    set_east_asia(run)
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(90, 90, 90)
                document.add_section(WD_SECTION.NEW_PAGE)
                first_title = False
            else:
                document.add_heading(title, level=level)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.match(
                r"^\|?\s*:?-{3,}", lines[index + 1]):
            rows = []
            rows.append([part.strip() for part in line.strip("|").split("|")])
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([part.strip() for part in lines[index].strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue

        bullet = re.match(r"^-\s+(.*)$", line)
        if bullet:
            value = plain(bullet.group(1))
            if value.startswith("[ ]"):
                value = "☐ " + value[3:].strip()
            paragraph = document.add_paragraph(value, style="List Bullet")
            index += 1
        elif re.match(r"^\d+\.\s+", line):
            value = plain(re.sub(r"^\d+\.\s+", "", line))
            paragraph = document.add_paragraph(value, style="List Number")
            index += 1
        else:
            parts = [line]
            index += 1
            while index < len(lines):
                nxt = lines[index].rstrip()
                if (not nxt or nxt.startswith(("#", "```", "|", "- "))
                        or re.match(r"^\d+\.\s+", nxt)):
                    break
                parts.append(nxt)
                index += 1
            paragraph = document.add_paragraph(plain(" ".join(parts)))
        for run in paragraph.runs:
            set_east_asia(run)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "docs" / "TECH_TRANSFER.md")
    parser.add_argument(
        "--output", type=Path,
        default=ROOT / "docs" / "인지검사_문항생성_파이프라인_기술이전서_v1.0.docx",
    )
    args = parser.parse_args()
    convert(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
